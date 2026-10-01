#!/usr/bin/env python3
"""Planilha de auditoria: os top-5 do BGE+reranker, para um especialista julgar.

PARA QUE SERVE. Antes de investir na geração, é preciso saber se o contexto que
a recuperação entrega dá para responder. O Strict Recall@5 responde isso contra
o GABARITO — e o gabarito é derivado de citação literal, escrita por LLM e
conferida à mão. Ele pode estar incompleto (norma cita norma: medido em 30/09,
2 casos em 9 lidos eram falso negativo da métrica). Só leitura humana do que o
sistema DE FATO devolve fecha essa lacuna.

TRÊS DECISÕES DE DESENHO, porque planilha mal montada produz número pior que
nenhum:

  A ORDEM É EMBARALHADA. O juiz vê os 5 trechos em ordem aleatória, não na
  ordem do ranking. Se visse, a posição 1 viraria âncora — "o sistema pôs em
  primeiro, deve ser bom". A posição verdadeira fica só na chave.

  NADA DO MODELO NA PLANILHA DO JUIZ. Sem nota do reranker, sem marcação de
  qual trecho está no gabarito, sem a resposta-referência. Se o juiz vê a
  resposta esperada, ele procura confirmação em vez de julgar; a auditoria
  passa a medir concordância com o gabarito, que é exatamente o que ela deveria
  checar de fora.

  O TIPO DA PERGUNTA NÃO APARECE. Saber que a pergunta é "comparativa" faria o
  juiz procurar duas normas antes de ler; ele não precisa do rótulo para julgar
  se o trecho responde.

SÓ PERGUNTAS RESPONDÍVEIS (decidido em 30/09/2026). A primeira versão incluía
irrespondíveis como controle, e estava errada: para dizer "estes 5 trechos não
respondem porque o acervo não tem a resposta", o juiz teria de garantir que a
resposta não está em nenhuma das 54 normas — o que ele não pode verificar lendo
5 trechos. Ele responderia "não" por padrão, sem informação, ou "sim" por
engano num trecho só topicamente parecido. A abstenção já é validada por número,
em metades separadas, contra 125 irrespondíveis escritas de propósito; não
precisa de juiz. O que precisa de juiz é o que esta planilha pergunta: para uma
pergunta que TEM resposta no acervo, o contexto recuperado basta?

A AMOSTRA é estratificada por tipo × acerto: metade das perguntas sorteadas o
sistema acerta e metade erra. Sortear só acertos mediria satisfação, não
qualidade. E vem em LOTES independentemente estratificados: o lote 1 já é uma
amostra válida, então dá para parar depois dele se o tempo acabar.

Uso:
  python eval/experimento_embedding/montar_auditoria.py
  python eval/experimento_embedding/montar_auditoria.py --por-celula 5 --lotes 3
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
import sys
from collections import defaultdict
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]

# A vista BALANCEADA 75/75/75, que é a população do artigo desde 30/09/2026.
GABARITO = RAIZ / "eval/data/golden_dezembro.jsonl"
R = RAIZ / "eval/results/retrieval/rankings"
SCORES = RAIZ / "eval/results/indexes/rerank_scores_bge_reranker_v2_m3.json"
CORPUS = RAIZ / "data/processed/documents.jsonl"
SAIDA = AQUI / "auditoria"

K = 5
SEMENTE = 2026

for _f in (sys.stdout, sys.stderr):
    try:
        _f.reconfigure(encoding="utf-8")
    except Exception:
        pass


def limpar(t: str) -> str:
    """Colapsa espaço em branco. Texto normativo vem com quebra de linha do PDF,
    e célula de planilha com quebra fica ilegível."""
    return re.sub(r"\s+", " ", (t or "").strip())


def escrever(caminho: Path, cabecalho: list[str], linhas: list[list]) -> None:
    # utf-8-sig e separador ';': é o que o Excel em português abre com as
    # colunas já separadas, sem passar pelo assistente de importação.
    with caminho.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_ALL)
        w.writerow(cabecalho)
        w.writerows(linhas)
    print(f"  ✓ {caminho.relative_to(RAIZ)}  ({len(linhas)} linhas)")


# Uma planilha só, em blocos de 5 linhas por pergunta: o juiz lê a pergunta uma
# vez, julga os 5 trechos ao lado, e responde "os 5 bastam?" na mesma linha em
# que a pergunta começa. Dois CSVs separados obrigavam a casar `qid` entre
# arquivos à mão — é onde a atenção se perde e o julgamento piora.
CAB_XLSX = [
    ("lote", 6), ("qid", 8), ("pergunta", 46), ("trecho", 7),
    ("documento", 26), ("título", 26), ("pág.", 6),
    ("texto do trecho", 88),
    ("JUIZ: relevância 0-3", 13), ("JUIZ: observação do trecho", 26),
    ("JUIZ: os 5 bastam? S/PARCIAL/N", 15), ("JUIZ: o que falta", 30),
]


def escrever_xlsx(caminho: Path, blocos: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "auditoria"

    cinza = PatternFill("solid", fgColor="EFEFEF")
    amarelo = PatternFill("solid", fgColor="FFF4CC")
    topo = Border(top=Side(style="medium", color="999999"))
    alto = Alignment(vertical="top", wrap_text=True)
    centro = Alignment(vertical="center", horizontal="center", wrap_text=True)

    for i, (nome, largura) in enumerate(CAB_XLSX, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="444444")
        c.alignment = centro
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "D2"
    ws.row_dimensions[1].height = 32

    linha = 2
    for n, b in enumerate(blocos):
        ini = linha
        for pos, t in enumerate(b["trechos"], start=1):
            vals = [b["lote"], b["qid"], b["pergunta"] if pos == 1 else None,
                    pos, t["documento"], t["titulo"], t["pagina"], t["texto"],
                    None, None, None, None]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=linha, column=col, value=v)
                c.alignment = alto
                if col in (9, 10):
                    c.fill = amarelo
                elif n % 2 == 0:
                    c.fill = cinza
                if linha == ini:
                    c.border = topo
            linha += 1
        fim = linha - 1
        # a pergunta e o veredito valem para o bloco inteiro
        for col in (1, 2, 3, 11, 12):
            ws.merge_cells(start_row=ini, start_column=col,
                           end_row=fim, end_column=col)
            c = ws.cell(row=ini, column=col)
            c.alignment = centro if col in (1, 2) else alto
            if col in (11, 12):
                c.fill = amarelo

    rel = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True,
                         showErrorMessage=True,
                         error="Use 0, 1, 2 ou 3.", errorTitle="Nota inválida")
    suf = DataValidation(type="list", formula1='"S,PARCIAL,N"',
                         allow_blank=True, showErrorMessage=True,
                         error="Use S, PARCIAL ou N.", errorTitle="Valor inválido")
    ws.add_data_validation(rel)
    ws.add_data_validation(suf)
    rel.add(f"I2:I{linha - 1}")
    suf.add(f"K2:K{linha - 1}")

    guia = wb.create_sheet("como preencher")
    guia.column_dimensions["A"].width = 112
    for i, par in enumerate(GUIA_XLSX, start=1):
        texto, negrito = par
        c = guia.cell(row=i, column=1, value=texto)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.font = Font(bold=negrito, size=13 if negrito else 11)

    wb.save(caminho)
    print(f"  ✓ {caminho.relative_to(RAIZ)}  ({len(blocos)} perguntas, "
          f"{linha - 2} linhas de trecho)")


GUIA_XLSX = [
    ("Auditoria da recuperação — como preencher", True),
    ("", False),
    ("Uma pergunta entra; o sistema busca em 54 normas (brasileiras e "
     "internacionais, em português e inglês) e devolve os 5 trechos que "
     "considerou mais relevantes. A resposta ao usuário é gerada SÓ com esses "
     "5 trechos. Estamos auditando a busca, antes de construir o gerador.", False),
    ("", False),
    ("Na aba 'auditoria', cada pergunta ocupa um bloco de 5 linhas. As colunas "
     "amarelas são suas.", False),
    ("", False),
    ("1) Para cada trecho, preencha 'relevância 0-3'", True),
    ("0 = não tem relação com a pergunta", False),
    ("1 = mesmo assunto, mas não contribui para a resposta", False),
    ("2 = contém PARTE do que a resposta precisa", False),
    ("3 = contém, sozinho, o suficiente para responder", False),
    ("", False),
    ("Julgue cada trecho CONTRA A PERGUNTA, não contra os outros trechos. Não "
     "há cota: os 5 podem ser 0, ou os 5 podem ser 3.", False),
    ("A ORDEM DOS TRECHOS É ALEATÓRIA. A coluna 'trecho' é posição na planilha, "
     "não a posição que o sistema deu — não tente inferir qual ele achou melhor.", True),
    ("", False),
    ("2) Depois dos 5, preencha 'os 5 bastam?' na primeira linha do bloco", True),
    ("S = sim, dá para responder a pergunta de forma completa e correta", False),
    ("PARCIAL = dá para responder parte, ou responder com ressalva", False),
    ("N = não dá", False),
    ("", False),
    ("Se marcar PARCIAL ou N, diga em 'o que falta' o que faltou: o "
     "dispositivo, a norma, o lado da comparação, o número, o prazo.", False),
    ("", False),
    ("Duas coisas que evitam mal-entendido", True),
    ("TODAS as perguntas aqui têm resposta no acervo. Você não precisa decidir "
     "se a resposta existe em algum lugar — só se ela está nestes 5 trechos. "
     "Marcar N significa 'a busca não trouxe', nunca 'não existe'.", False),
    ("Metade das perguntas sorteadas são casos em que o sistema, pelo gabarito "
     "automático, ERROU. Isso é de propósito: auditar só os acertos mediria "
     "satisfação, não qualidade. Não se surpreenda se vários conjuntos "
     "parecerem ruins.", False),
    ("", False),
    ("Os enunciados estão em português; os trechos podem estar em português ou "
     "inglês, porque o acervo é bilíngue. Isso é normal, não é defeito.", False),
    ("", False),
    ("Faça o LOTE 1 inteiro primeiro — ele já é uma amostra válida.", True),
]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--por-celula", type=int, default=5,
                    help="perguntas por (tipo × acerto) por lote; padrão 5")
    ap.add_argument("--lotes", type=int, default=3)
    args = ap.parse_args()

    gold = {g["qid"]: g for g in
            (json.loads(l) for l in GABARITO.open(encoding="utf-8") if l.strip())}
    bge = json.loads((R / "B_dense_bge_m3.json").read_text(encoding="utf-8"))
    rr = json.loads((R / "B_bge_m3_rerank.json").read_text(encoding="utf-8"))
    notas = json.loads(SCORES.read_text(encoding="utf-8"))["scores"]

    trechos = {}
    with CORPUS.open(encoding="utf-8") as fh:
        for linha in fh:
            if linha.strip():
                r = json.loads(linha)
                trechos[r["id"]] = r

    def top5(q: str) -> list[str]:
        """O top-5 do BGE+reranker, exatamente como o sistema devolveria."""
        if q in rr:
            return rr[q][:K]
        pool = bge.get(q, [])[:100]
        nq = notas.get(q, {})
        return sorted(pool, key=lambda c: -nq.get(c, -1e9))[:K]

    def acerta(q: str) -> bool:
        gr = gold[q]["qrels_grupos"]
        if not gr:
            return False
        topo = set(top5(q))
        return all(any(c in topo for c in g) for g in gr)

    # ── amostra estratificada por tipo × acerto, só respondíveis ────────────
    celulas: dict[tuple, list[str]] = defaultdict(list)
    for q, g in gold.items():
        if q not in bge or not notas.get(q) or not g["qrels_grupos"]:
            continue
        celulas[(g["question_type"], "acerta" if acerta(q) else "erra")].append(q)

    rnd = random.Random(SEMENTE)
    for v in celulas.values():
        v.sort()
        rnd.shuffle(v)

    print("células disponíveis (tipo × acerto no top-5):")
    for k in sorted(celulas):
        print(f"  {k[0]:14s} {k[1]:8s} {len(celulas[k]):4d}")

    cota = {k: args.por_celula for k in celulas}

    escolhidas: list[tuple[int, str]] = []
    cursor = {k: 0 for k in celulas}
    for lote in range(1, args.lotes + 1):
        for k in sorted(celulas):
            pega = celulas[k][cursor[k]:cursor[k] + cota[k]]
            cursor[k] += len(pega)
            escolhidas.extend((lote, q) for q in pega)

    print(f"\namostra: {len(escolhidas)} perguntas em {args.lotes} lotes "
          f"({len(escolhidas)*K} julgamentos de trecho)")
    for lote in range(1, args.lotes + 1):
        n = sum(1 for l, _ in escolhidas if l == lote)
        print(f"  lote {lote}: {n} perguntas, {n*K} trechos")

    SAIDA.mkdir(exist_ok=True)

    # ── planilha 1: um trecho por linha ─────────────────────────────────────
    linhas_t, linhas_p, linhas_c, blocos = [], [], [], []
    for lote, q in escolhidas:
        g = gold[q]
        ordem_real = top5(q)
        exibicao = list(enumerate(ordem_real, start=1))   # (rank real, cid)
        rnd.shuffle(exibicao)
        bloco = {"lote": lote, "qid": q, "pergunta": limpar(g["question"]),
                 "trechos": []}

        for pos_exib, (rank_real, cid) in enumerate(exibicao, start=1):
            t = trechos.get(cid, {})
            md = t.get("metadata", {})
            bloco["trechos"].append({
                "documento": md.get("document_id", ""),
                "titulo": limpar(md.get("title", "")),
                "pagina": md.get("page", ""),
                "texto": limpar(t.get("text", "")),
            })
            linhas_t.append([
                lote, q, limpar(g["question"]), pos_exib,
                cid,
                md.get("document_id", ""),
                limpar(md.get("title", "")),
                md.get("page", ""),
                limpar(md.get("section_path", "")),
                limpar(t.get("text", "")),
                "", "",                      # colunas do juiz
            ])
            anotado = any(cid in grupo for grupo in g["qrels_grupos"])
            linhas_c.append([
                lote, q, pos_exib, rank_real, cid,
                f"{notas.get(q, {}).get(cid, float('nan')):.4f}",
                "SIM" if anotado else "não",
            ])

        linhas_p.append([
            lote, q, limpar(g["question"]),
            "", "", "",                      # colunas do juiz
        ])
        linhas_c.append([
            lote, q, "", "", "RESPOSTA-REFERÊNCIA",
            g["question_type"],
            limpar(g.get("reference_answer") or ""),
        ])
        blocos.append(bloco)

    print()
    escrever_xlsx(SAIDA / "AUDITORIA_RECUPERACAO.xlsx", blocos)
    escrever(SAIDA / "auditoria_trechos.csv", [
        "lote", "qid", "pergunta", "trecho_n", "trecho_id",
        "documento", "titulo", "pagina", "secao", "texto_do_trecho",
        "JUIZ_relevancia_0a3", "JUIZ_observacao",
    ], linhas_t)

    escrever(SAIDA / "auditoria_perguntas.csv", [
        "lote", "qid", "pergunta",
        "JUIZ_suficiente_S_PARCIAL_N", "JUIZ_o_que_falta", "JUIZ_observacao",
    ], linhas_p)

    escrever(SAIDA / "auditoria_CHAVE_nao_abrir_antes.csv", [
        "lote", "qid", "trecho_n_exibido", "rank_real_no_ranking", "trecho_id",
        "nota_reranker", "esta_no_gabarito",
    ], linhas_c)

    (SAIDA / "auditoria_INSTRUCOES.md").write_text(INSTRUCOES, encoding="utf-8")
    print(f"  ✓ {(SAIDA / 'auditoria_INSTRUCOES.md').relative_to(RAIZ)}")
    print("\nMANDE SÓ o AUDITORIA_RECUPERACAO.xlsx — pergunta e trechos na mesma")
    print("aba, com as colunas do juiz em amarelo e listas suspensas. Os CSVs")
    print("ficam para quem preferir processar em texto.")
    print("A CHAVE NÃO vai para o especialista: ela contém o gabarito.")


INSTRUCOES = """# Auditoria da recuperação — instruções

Obrigado por auditar. São duas planilhas e o trabalho é em duas passadas por
pergunta. Não abra `auditoria_CHAVE_nao_abrir_antes.csv`: ela tem o gabarito, e
consultá-la antes invalida o julgamento.

## O que o sistema faz

Uma pergunta entra; o sistema busca em 54 normas (brasileiras e internacionais,
em português e inglês) e devolve os **5 trechos** que considerou mais relevantes.
Cada trecho tem cerca de 1.200 caracteres. A pergunta é respondida **só** com
esses 5 trechos — o gerador não tem acesso a mais nada.

Estamos auditando a etapa de busca, antes de construir o gerador. A pergunta de
fundo é simples: **esses 5 trechos bastam?**

## Passada 1 — `auditoria_trechos.csv`, trecho por trecho

Cada linha é um trecho. Preencha `JUIZ_relevancia_0a3`:

| nota | quando |
|---|---|
| **0** | não tem relação com a pergunta |
| **1** | mesmo assunto, mas não contribui para a resposta |
| **2** | contém **parte** do que a resposta precisa |
| **3** | contém, sozinho, o suficiente para responder |

Duas observações que mudam o resultado:

- Julgue o trecho **contra a pergunta**, não contra os outros trechos. Não há
  cota: os 5 podem ser 0, ou os 5 podem ser 3.
- **A ordem na planilha é aleatória.** `trecho_n` é posição de exibição, não a
  posição que o sistema deu. Não tente inferir qual o sistema achou melhor.

Use `JUIZ_observacao` quando a nota não contar a história — por exemplo "é o
artigo certo mas a frase que responde está no trecho anterior", que é o defeito
mais comum que encontramos.

## Passada 2 — `auditoria_perguntas.csv`, a pergunta inteira

Depois de julgar os 5 trechos de uma pergunta, responda: **com esses 5 trechos,
é possível responder a pergunta de forma completa e correta?**

| valor | quando |
|---|---|
| **S** | sim, completa e correta |
| **PARCIAL** | dá para responder parte, ou responder com ressalva |
| **N** | não dá |

Se for `PARCIAL` ou `N`, diga em `JUIZ_o_que_falta` o que faltou: o dispositivo,
a norma, o lado da comparação, o número, o prazo.

**Todas as perguntas desta planilha têm resposta no acervo.** Você não precisa
decidir se a resposta existe em algum lugar — só se ela está nestes 5 trechos.
Quando marcar `N`, está dizendo "a busca não trouxe", não "não existe".

Os enunciados estão em português. Os trechos podem estar em português ou em
inglês, porque o acervo é bilíngue — isso é normal e não é defeito.

## Ordem sugerida

Faça o **lote 1** inteiro primeiro. Ele já é uma amostra estratificada válida, e
com ele fechamos uma primeira leitura. Os lotes 2 e 3 refinam.

## O que vai ser feito com isso

Três coisas: (1) medir concordância entre o seu julgamento e o gabarito
automático, que é derivado de citação literal e pode estar incompleto;
(2) encontrar trechos que respondem e não estão no gabarito — já achamos dois
casos, em que uma norma reafirma a regra de outra; (3) decidir se a recuperação
está boa o bastante para passarmos à geração.

Metade das perguntas sorteadas são casos em que o sistema, pelo gabarito
automático, **errou**. Isso é de propósito: auditar só os acertos mediria
satisfação, não qualidade. Não se surpreenda se vários conjuntos de 5 trechos
parecerem ruins.
"""


if __name__ == "__main__":
    main()
